"""
=========================================================
KiG POS
=========================================================

Datei:
    berichte/excel_layout.py

Beschreibung:
    Gemeinsames Aussehen aller Excel-Ausgaben.

    Früher schrieb jede Ausgabe ihre Zeilen selbst: eine fette
    Titelzeile, eine fette Kopfzeile, sonst nichts. Ausgedruckt
    sah ein Kassenbuch aus wie ein Datenabzug.

    Jetzt beginnt jedes Blatt gleich:

        [KiG-Logo]
        Kassenbuch September 2026                  (groß)
        erstellt am 13.09.2026 12:40                (grau)
        ───────────────────────────────────────── (orange)

    Darunter folgen Abschnitte mit Überschrift und Tabellen mit
    orangefarbener Kopfzeile, abwechselnd hinterlegten Zeilen,
    Beträgen als Währung und einer fetten Summenzeile. Beim
    Drucken passt das Blatt in die Seitenbreite, die Kopfzeile
    der Tabelle wiederholt sich auf jeder Seite, und unten steht
    "Seite x von y".

    Verwendung:

        blatt = Exportblatt(mappe.active, "Kassenbuch", "September 2026",
                            breiten=(12, 14, 14))
        blatt.ueberschrift("Einträge")
        blatt.tabelle(("Datum", "Betrag"), zeilen, formate=("text", "geld"))
        blatt.drucken()

Version:
    1.0.0
=========================================================
"""

from datetime import datetime
from io import BytesIO

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config


# Papierfarben - unabhängig davon, ob das Programm hell oder dunkel
# läuft (siehe berichte/__init__.py).
ORANGE = "F44611"
ORANGE_HELL = "FDECE5"
TEXT = "222222"
GRAU = "6B6B6B"
LINIE = "D9D9D9"
ZEBRA = "F7F7F7"
WEISS = "FFFFFF"
ROT = "C62828"
GRUEN = "2E7D32"

GELD = '#,##0.00 "€";[Red]-#,##0.00 "€"'
# "General": 442 bleibt 442 und 0,5 bleibt 0,5 - ein festes Format
# haengte an ganze Zahlen ein Komma ("442,").
ZAHL = "General"
PROZENT = '0.0 %'

LOGO_HOEHE_PX = 62


def erstellt_am():
    """"erstellt am 13.09.2026 um 12:40"."""

    return f"erstellt am {datetime.now():%d.%m.%Y} um {datetime.now():%H:%M}"


def logo_bild(hoehe_px=LOGO_HOEHE_PX):
    """Das Vereinslogo als Bild für openpyxl - oder None, wenn es
    fehlt oder Pillow nicht da ist.

    Verkleinert wird vorher: Das Original ist 1944 Bildpunkte breit
    und machte jede Mappe um 200 KB schwerer.
    """

    try:
        from openpyxl.drawing.image import Image as ExcelBild
        from PIL import Image
    except ImportError:
        return None

    pfad = getattr(config, "LOGO_PATH", None)

    if not pfad or not pfad.exists():
        return None

    try:
        with Image.open(pfad) as original:
            breite = round(original.width * hoehe_px / original.height)
            klein = original.convert("RGBA").resize(
                (breite * 2, hoehe_px * 2), Image.LANCZOS
            )

        puffer = BytesIO()
        klein.save(puffer, format="PNG")
        puffer.seek(0)

        # Über den Puffer geöffnet, damit openpyxl ein Bild mit
        # bekanntem Format bekommt (ein verkleinertes hat keins).
        bild = ExcelBild(Image.open(puffer))

    except (OSError, ValueError):
        return None

    # Doppelt so fein gerechnet, in halber Größe gezeigt - gedruckt
    # bleibt die Schrift im Logo scharf.
    bild.width = breite
    bild.height = hoehe_px

    return bild


class Exportblatt:
    """Ein Tabellenblatt im Aussehen der KiG-Ausgaben.

    blatt       das openpyxl-Arbeitsblatt
    titel       große Überschrift ("Kassenbuch September 2026")
    untertitel  wofür die Zahlen gelten; "erstellt am ..." kommt
                immer dazu
    breiten     Spaltenbreiten in Zeichen - ihre Anzahl bestimmt
                auch, wie weit Linien und Überschriften reichen
    quer        Querformat beim Drucken
    """

    def __init__(self, blatt, titel, untertitel="", breiten=(20,) * 6,
                 quer=True):

        self.blatt = blatt
        self.titel = titel
        self.spalten = len(breiten)
        self.quer = quer

        # Die nächste freie Zeile
        self.zeile = 1

        # Zeile der letzten Tabellenkopfzeile - sie wiederholt sich
        # beim Drucken auf jeder Seite.
        self.kopfzeile = None

        blatt.sheet_view.showGridLines = False

        for nummer, breite in enumerate(breiten, start=1):
            blatt.column_dimensions[get_column_letter(nummer)].width = breite

        self._kopf(titel, untertitel)

    # =====================================================
    # Kopf
    # =====================================================

    def _kopf(self, titel, untertitel):

        blatt = self.blatt

        # Zeile 1: das Logo. 50 Punkt sind gut 66 Bildpunkte - das
        # Logo passt hinein, ohne in den Titel zu ragen.
        blatt.row_dimensions[1].height = 50

        bild = logo_bild()

        if bild is not None:
            blatt.add_image(bild, "A1")

        # Rechts oben der Absender
        rechts = blatt.cell(row=1, column=self.spalten, value="KiG POS")
        rechts.font = Font(bold=True, size=11, color=ORANGE)
        rechts.alignment = Alignment(horizontal="right", vertical="bottom")

        zelle = blatt.cell(row=2, column=1, value=titel)
        zelle.font = Font(bold=True, size=18, color=TEXT)
        blatt.row_dimensions[2].height = 28

        teile = [teil for teil in (untertitel, erstellt_am()) if teil]

        zelle = blatt.cell(row=3, column=1, value="   ·   ".join(teile))
        zelle.font = Font(size=10, color=GRAU)

        # Zeile 4: orangefarbene Linie über die ganze Breite
        blatt.row_dimensions[4].height = 6

        for spalte in range(1, self.spalten + 1):
            blatt.cell(row=4, column=spalte).border = Border(
                bottom=Side(style="medium", color=ORANGE)
            )

        self.zeile = 6

    # =====================================================
    # Bausteine
    # =====================================================

    def leerzeile(self, anzahl=1):

        self.zeile += anzahl

    def ueberschrift(self, text):
        """Abschnittsüberschrift: orange, fett, mit feiner Linie."""

        if self.zeile > 6:
            self.zeile += 1

        zelle = self.blatt.cell(row=self.zeile, column=1, value=text)
        zelle.font = Font(bold=True, size=13, color=ORANGE)

        self.blatt.row_dimensions[self.zeile].height = 22

        for spalte in range(1, self.spalten + 1):
            self.blatt.cell(row=self.zeile, column=spalte).border = Border(
                bottom=Side(style="thin", color=ORANGE_HELL)
            )

        self.zeile += 2

    def text(self, inhalt, grau=False, fett=False, farbe=None):
        """Eine Zeile Fließtext über die ganze Breite."""

        zelle = self.blatt.cell(row=self.zeile, column=1, value=inhalt)
        zelle.font = Font(
            size=10, bold=fett, color=farbe or (GRAU if grau else TEXT)
        )

        self.zeile += 1

        return zelle

    def kennzahlen(self, paare, formate=None):
        """Bezeichnung und Wert untereinander - für "Einnahmen 120,00 €".

        paare:   ((Bezeichnung, Wert), ...)
        formate: je Paar "geld", "zahl" oder "text" (Vorgabe: geld)
        """

        formate = formate or ("geld",) * len(paare)

        for (bezeichnung, wert), format_ in zip(paare, formate):

            links = self.blatt.cell(row=self.zeile, column=1, value=bezeichnung)
            links.font = Font(size=11, color=GRAU)
            links.alignment = Alignment(vertical="center")

            rechts = self.blatt.cell(row=self.zeile, column=2, value=wert)
            rechts.font = Font(size=12, bold=True, color=TEXT)
            rechts.alignment = Alignment(horizontal="right", vertical="center")
            self._format(rechts, format_)

            for spalte in (1, 2):
                self.blatt.cell(row=self.zeile, column=spalte).border = Border(
                    bottom=Side(style="thin", color=LINIE)
                )

            self.blatt.row_dimensions[self.zeile].height = 20

            self.zeile += 1

    def tabelle(self, kopf, zeilen, formate=None, summe=None,
                umbrechen=(), hervorheben=None, filter_setzen=False):
        """Eine Tabelle mit Kopfzeile.

        kopf         Spaltenüberschriften
        zeilen       Werte je Zeile
        formate      je Spalte "text", "geld", "zahl", "prozent"
                     (Vorgabe: text)
        summe        optional eine Summenzeile (fett, Linie darüber)
        umbrechen    Spaltennummern (ab 0), deren Text umbricht
        hervorheben  optional f(zeile) -> Farbe (Hex) für die Schrift,
                     z. B. rot für auffällige Zeilen
        filter_setzen Excel-Filter auf die Kopfzeile

        Liefert (erste, letzte) Zeilennummer der Daten - für
        Diagramme, die auf die Tabelle verweisen.
        """

        blatt = self.blatt
        formate = formate or ("text",) * len(kopf)

        kopf_nr = self.zeile
        self.kopfzeile = kopf_nr

        duenn = Side(style="thin", color=LINIE)

        for spalte, text in enumerate(kopf, start=1):

            zelle = blatt.cell(row=kopf_nr, column=spalte, value=text)
            zelle.font = Font(bold=True, color=WEISS, size=10)
            zelle.fill = PatternFill("solid", fgColor=ORANGE)
            zelle.alignment = Alignment(
                horizontal="right" if formate[spalte - 1] in ("geld", "zahl", "prozent")
                else "left",
                vertical="center",
                wrap_text=True,
                indent=0 if formate[spalte - 1] in ("geld", "zahl", "prozent") else 1,
            )

        blatt.row_dimensions[kopf_nr].height = 22

        erste = kopf_nr + 1
        nummer = erste

        for position, werte in enumerate(zeilen):

            farbe = hervorheben(werte) if callable(hervorheben) else None

            for spalte, wert in enumerate(werte, start=1):

                zelle = blatt.cell(row=nummer, column=spalte, value=wert)
                zelle.font = Font(size=10, color=farbe or TEXT)
                zelle.border = Border(bottom=duenn)

                if position % 2 == 1:
                    zelle.fill = PatternFill("solid", fgColor=ZEBRA)

                format_ = formate[spalte - 1]

                zelle.alignment = Alignment(
                    horizontal="right" if format_ in ("geld", "zahl", "prozent")
                    else "left",
                    vertical="top",
                    wrap_text=(spalte - 1) in umbrechen,
                    indent=0 if format_ in ("geld", "zahl", "prozent") else 1,
                )

                self._format(zelle, format_)

            nummer += 1

        letzte = nummer - 1

        if filter_setzen and letzte >= erste:
            blatt.auto_filter.ref = (
                f"A{kopf_nr}:{get_column_letter(len(kopf))}{letzte}"
            )

        if summe is not None:

            for spalte, wert in enumerate(summe, start=1):

                zelle = blatt.cell(row=nummer, column=spalte, value=wert)
                zelle.font = Font(bold=True, size=10, color=TEXT)
                zelle.fill = PatternFill("solid", fgColor=ORANGE_HELL)
                zelle.border = Border(top=Side(style="medium", color=ORANGE))

                format_ = formate[spalte - 1]

                zelle.alignment = Alignment(
                    horizontal="right" if format_ in ("geld", "zahl", "prozent")
                    else "left",
                    indent=0 if format_ in ("geld", "zahl", "prozent") else 1,
                )

                self._format(zelle, format_)

            nummer += 1

        self.zeile = nummer + 1

        return erste, letzte

    @staticmethod
    def _format(zelle, format_):

        if format_ == "geld":
            zelle.number_format = GELD
        elif format_ == "zahl":
            zelle.number_format = ZAHL
        elif format_ == "prozent":
            zelle.number_format = PROZENT

    # =====================================================
    # Drucken
    # =====================================================

    def drucken(self, titel_wiederholen=True):
        """Seitenbreite, Ränder, Fußzeile - und die Kopfzeile der
        (letzten) Tabelle auf jeder Seite."""

        blatt = self.blatt

        blatt.page_setup.orientation = "landscape" if self.quer else "portrait"
        blatt.page_setup.paperSize = blatt.PAPERSIZE_A4
        blatt.page_setup.fitToWidth = 1
        blatt.page_setup.fitToHeight = 0
        blatt.sheet_properties.pageSetUpPr.fitToPage = True

        blatt.page_margins.left = 0.5
        blatt.page_margins.right = 0.5
        blatt.page_margins.top = 0.6
        blatt.page_margins.bottom = 0.7

        blatt.print_options.horizontalCentered = True

        blatt.oddFooter.left.text = "KiG POS"
        blatt.oddFooter.left.size = 8
        blatt.oddFooter.center.text = self.titel
        blatt.oddFooter.center.size = 8
        blatt.oddFooter.right.text = "Seite &P von &N"
        blatt.oddFooter.right.size = 8

        if titel_wiederholen and self.kopfzeile:
            blatt.print_title_rows = f"{self.kopfzeile}:{self.kopfzeile}"

            # Beim Blättern am Bildschirm bleibt sie ebenso stehen.
            blatt.freeze_panes = f"A{self.kopfzeile + 1}"


def dateiname_sicher(text, ersatz="liste"):
    """Aus "Sommerfest 2026!" wird "Sommerfest_2026_"."""

    sicher = "".join(
        zeichen if zeichen.isalnum() or zeichen in " -_" else "_"
        for zeichen in (text or "")
    ).strip().replace(" ", "_")

    return sicher or ersatz


def blattname(text, ersatz="Blatt"):
    """Excel erlaubt 31 Zeichen und kein : \\ / ? * [ ]."""

    sauber = "".join(
        zeichen for zeichen in (text or "") if zeichen not in ':\\/?*[]'
    ).strip()

    return (sauber or ersatz)[:31]
