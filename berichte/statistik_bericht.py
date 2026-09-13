"""
=========================================================
KiG POS
=========================================================

Datei:
    berichte/statistik_bericht.py

Beschreibung:
    Die Verkaufsstatistik zum Ausdrucken - als Excel-Mappe und
    als PDF, beide für dieselbe Auswahl (Zeitraum, Event,
    Kategorie) wie auf dem Bildschirm.

    Excel, Blatt "Zusammenfassung":
        Kennzahlen, Umsatz nach Kategorie mit Kreisdiagramm,
        Verkäufe je Artikel mit Balkendiagramm - die Farben sind
        die der Kategorien, wie in der App.
    Excel, Blatt "Einzelverkäufe":
        jede verkaufte Einheit, mit Filter in der Kopfzeile.

    PDF:
        Kennzahlen als Kacheln, Kreis mit Legende, Balken je
        Artikel und die Tabelle je Artikel. Die Einzelverkäufe
        stehen nur in Excel - als PDF wären es schnell zwanzig
        Seiten, die niemand liest.

    Die Daten holt sammeln() einmal aus der Datenbank; excel()
    und pdf() schreiben nur noch.

Version:
    1.0.0
=========================================================
"""

from datetime import datetime

import geldformat

from berichte.excel_layout import Exportblatt, ORANGE
from berichte.pdf_bericht import PdfBericht, ROT


# Ersatzfarben wie im Kreisdiagramm der App (category_pie.py) - eine
# Kategorie ohne eigene Farbe soll im Bericht dieselbe haben wie am
# Bildschirm.
FALLBACK_FARBEN = (
    "#1976D2", "#F57C00", "#43A047", "#7B1FA2",
    "#00838F", "#C62828", "#5D4037", "#616161",
)


def _farbe(wert, position):

    wert = (wert or "").strip()

    if len(wert.lstrip("#")) >= 6:
        try:
            int(wert.lstrip("#")[:6], 16)
            return "#" + wert.lstrip("#")[:6].upper()
        except ValueError:
            pass

    return FALLBACK_FARBEN[position % len(FALLBACK_FARBEN)]


def _datum(iso):

    try:
        return datetime.strptime(iso, "%Y-%m-%d").strftime("%d.%m.%Y")
    except (TypeError, ValueError):
        return str(iso or "-")


# =========================================================
# Daten
# =========================================================

def sammeln(db, auswahl, beschreibung):
    """Alles für den Bericht in einem Wörterbuch.

    auswahl:      (date_from, date_to, event_id, category_id)
    beschreibung: in Worten, wofür die Zahlen gelten
                  ("Sommerfest | Getränke | 01.09. - 13.09.2026")
    """

    kennzahlen = db.get_period_totals(*auswahl)
    artikel = db.get_article_sales(*auswahl)
    kategorien_roh = db.get_category_revenues(*auswahl)
    einzeln = db.get_statistic_sale_items(*auswahl)

    # Farbe je Kategorie einmal festlegen - Kreis, Balken und
    # Tabellen benutzen dieselbe.
    farben = {}
    kategorien = []

    for position, (name, betrag, farbe) in enumerate(kategorien_roh):
        farben[name] = _farbe(farbe, position)
        kategorien.append((name, betrag, farben[name]))

    for eintrag in artikel:
        eintrag["farbe_hex"] = farben.get(eintrag["kategorie"], FALLBACK_FARBEN[0])

    return {
        "beschreibung": beschreibung,
        "kennzahlen": kennzahlen,
        "artikel": artikel,
        "kategorien": kategorien,
        "einzeln": einzeln,
    }


def _kennzahlen_paare(kennzahlen):
    """(Bezeichnung, Wert, Format) - in der Reihenfolge des Berichts."""

    paare = [
        ("Einnahmen", kennzahlen["revenue"], "geld"),
        ("Ausgaben (Einkauf)", kennzahlen["expenses"], "geld"),
        ("Gewinn", kennzahlen["profit"], "geld"),
        ("Verkaufte Einheiten", kennzahlen["quantity"], "zahl"),
        ("Bons", kennzahlen["receipts"], "zahl"),
    ]

    if kennzahlen.get("kig_karte") or kennzahlen.get("gutschein"):
        paare += [
            ("Entwertet: KiG Karte", kennzahlen["kig_karte"], "geld"),
            ("Entwertet: Gutschein", kennzahlen["gutschein"], "geld"),
            ("Bar eingenommen", kennzahlen["bar"], "geld"),
        ]

    return paare


# =========================================================
# Excel
# =========================================================

def excel(daten, pfad):
    """Schreibt die Mappe nach `pfad` und liefert den Pfad."""

    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.series import DataPoint

    mappe = Workbook()

    # -----------------------------------------------------
    # Zusammenfassung
    # -----------------------------------------------------

    blatt = Exportblatt(
        mappe.active, "Verkaufsstatistik", daten["beschreibung"],
        breiten=(30, 18, 11, 15, 15, 3, 14, 14, 14, 14, 14, 14),
        quer=False,
    )
    blatt.blatt.title = "Zusammenfassung"

    kennzahlen = daten["kennzahlen"]
    paare = _kennzahlen_paare(kennzahlen)

    blatt.ueberschrift("Kennzahlen")
    blatt.kennzahlen(
        [(bezeichnung, round(wert or 0, 2)) for bezeichnung, wert, _f in paare],
        formate=[format_ for _b, _w, format_ in paare],
    )

    # ---- Kategorien mit Kreis --------------------------------

    kategorien = daten["kategorien"]
    gesamt = sum(betrag for _n, betrag, _f in kategorien) or 1

    blatt.ueberschrift("Umsatz nach Kategorie")

    kreis_zeile = blatt.zeile

    erste, letzte = blatt.tabelle(
        ("Kategorie", "Umsatz", "Anteil"),
        [
            (name, round(betrag, 2), betrag / gesamt)
            for name, betrag, _farbe in kategorien
        ],
        formate=("text", "geld", "prozent"),
        summe=("Summe", round(sum(b for _n, b, _f in kategorien), 2), 1 if kategorien else 0),
    )

    if kategorien and letzte >= erste:

        kreis = PieChart()
        # Ohne eigenen Titel - die Überschrift steht daneben.
        kreis.height = 7.5
        kreis.width = 11

        kreis.add_data(
            Reference(blatt.blatt, min_col=2, min_row=erste - 1, max_row=letzte),
            titles_from_data=True,
        )
        kreis.set_categories(
            Reference(blatt.blatt, min_col=1, min_row=erste, max_row=letzte)
        )

        reihe = kreis.series[0]

        for position, (_name, _betrag, farbe) in enumerate(kategorien):
            punkt = DataPoint(idx=position)
            punkt.graphicalProperties.solidFill = farbe.lstrip("#")
            punkt.graphicalProperties.line.solidFill = "FFFFFF"
            reihe.dPt.append(punkt)

        reihe.dLbls = DataLabelList()
        reihe.dLbls.showPercent = True
        reihe.dLbls.showVal = False
        reihe.dLbls.showCatName = False
        reihe.dLbls.showSerName = False
        reihe.dLbls.showLeaderLines = False

        blatt.blatt.add_chart(kreis, f"G{kreis_zeile}")

        # Der Kreis ist gut 15 Zeilen hoch - der nächste Abschnitt
        # soll nicht darunter liegen.
        blatt.zeile = max(blatt.zeile, kreis_zeile + 16)

    # ---- Artikel mit Balken ----------------------------------

    artikel = daten["artikel"]

    blatt.ueberschrift("Verkäufe je Artikel")

    balken_zeile = blatt.zeile

    erste, letzte = blatt.tabelle(
        ("Artikel", "Kategorie", "Menge", "Umsatz", "Gewinn"),
        [
            (e["name"], e["kategorie"], e["menge"], round(e["umsatz"], 2),
             round(e["gewinn"], 2))
            for e in artikel
        ],
        formate=("text", "text", "zahl", "geld", "geld"),
        summe=(
            "Summe", "",
            sum(e["menge"] for e in artikel),
            round(sum(e["umsatz"] for e in artikel), 2),
            round(sum(e["gewinn"] for e in artikel), 2),
        ),
    )

    if artikel and letzte >= erste:

        balken = BarChart()
        balken.type = "bar"
        balken.style = 10
        balken.legend = None
        # Seit openpyxl 3.1 sind Achsen ausgeblendet, solange man es
        # nicht anders sagt - dann fehlten die Artikelnamen.
        balken.x_axis.delete = False
        balken.y_axis.delete = False
        balken.y_axis.numFmt = '#,##0 "€"'
        balken.y_axis.majorGridlines = None

        # Größter Umsatz oben - wie in der App
        balken.x_axis.scaling.orientation = "maxMin"

        # ... und die Euro-Skala trotzdem unten statt oben
        balken.y_axis.crosses = "max"

        balken.width = 16
        balken.height = max(7.5, 0.6 * len(artikel) + 2)

        balken.add_data(
            Reference(blatt.blatt, min_col=4, min_row=erste - 1, max_row=letzte),
            titles_from_data=True,
        )
        balken.set_categories(
            Reference(blatt.blatt, min_col=1, min_row=erste, max_row=letzte)
        )

        reihe = balken.series[0]
        reihe.graphicalProperties.solidFill = ORANGE

        for position, eintrag in enumerate(artikel):
            punkt = DataPoint(idx=position)
            punkt.graphicalProperties.solidFill = eintrag["farbe_hex"].lstrip("#")
            punkt.graphicalProperties.line.solidFill = eintrag["farbe_hex"].lstrip("#")
            reihe.dPt.append(punkt)

        blatt.blatt.add_chart(balken, f"G{balken_zeile}")

    blatt.drucken(titel_wiederholen=False)

    # -----------------------------------------------------
    # Einzelverkäufe
    # -----------------------------------------------------

    einzeln = daten["einzeln"]

    zweites = Exportblatt(
        mappe.create_sheet("Einzelverkäufe"),
        "Einzelverkäufe", daten["beschreibung"],
        breiten=(22, 12, 16, 38, 8, 12, 12, 12),
    )

    zweites.tabelle(
        ("Event", "Datum", "Kategorie", "Artikel", "Menge", "Verkauf",
         "Einkauf", "Gewinn"),
        [
            (
                z["event_name"], _datum(z["business_date"]), z["category_name"],
                z["article_name"], z["quantity"], z["unit_price"],
                z["purchase_price"], z["profit"],
            )
            for z in einzeln
        ],
        formate=("text", "text", "text", "text", "zahl", "geld", "geld", "geld"),
        summe=(
            "Summe", "", "", "",
            sum(z["quantity"] for z in einzeln),
            round(sum(z["quantity"] * z["unit_price"] for z in einzeln), 2),
            round(sum(z["quantity"] * z["purchase_price"] for z in einzeln), 2),
            round(sum(z["profit"] for z in einzeln), 2),
        ),
        # Stornos rot, wie in der App
        hervorheben=lambda werte: "C62828" if (werte[4] or 0) < 0 else None,
        filter_setzen=True,
    )

    zweites.drucken()

    mappe.save(pfad)

    return pfad


# =========================================================
# PDF
# =========================================================

def pdf(daten, pfad):
    """Schreibt den Bericht als PDF nach `pfad` und liefert den Pfad."""

    geld = geldformat.geld

    bericht = PdfBericht("Verkaufsstatistik", daten["beschreibung"])

    kennzahlen = daten["kennzahlen"]

    bericht.ueberschrift("Kennzahlen")

    bericht.kennzahlen(
        [
            (bezeichnung, geld(wert) if format_ == "geld" else f"{wert:g}".replace(".", ","))
            for bezeichnung, wert, format_ in _kennzahlen_paare(kennzahlen)
        ],
        hervorheben=("Gewinn",),
    )

    if not daten["artikel"]:
        bericht.text("Im gewählten Zeitraum wurde nichts verkauft.")
        bericht.speichern(pfad)
        return pfad

    bericht.ueberschrift("Umsatz nach Kategorie", platz_danach=440)
    bericht.kreisdiagramm(daten["kategorien"], betrag_text=geld)

    artikel = daten["artikel"]

    bericht.ueberschrift("Umsatz je Artikel")
    bericht.balkendiagramm(
        [(e["name"], e["umsatz"], e["farbe_hex"]) for e in artikel],
        betrag_text=geld,
        zusatz_text=lambda i: f"{artikel[i]['kategorie']} · {artikel[i]['menge']:g} Stück",
    )

    bericht.ueberschrift("Verkäufe je Artikel")
    bericht.tabelle(
        ("Artikel", "Kategorie", "Menge", "Umsatz", "Gewinn"),
        [
            (e["name"], e["kategorie"], f"{e['menge']:g}", geld(e["umsatz"]),
             geld(e["gewinn"]))
            for e in artikel
        ],
        anteile=(0.34, 0.22, 0.12, 0.16, 0.16),
        ausrichtung=("left", "left", "right", "right", "right"),
        summe=(
            "Summe", "",
            f"{sum(e['menge'] for e in artikel):g}",
            geld(sum(e["umsatz"] for e in artikel)),
            geld(sum(e["gewinn"] for e in artikel)),
        ),
        farben=lambda werte: ROT if werte[4].startswith("-") else None,
    )

    bericht.speichern(pfad)

    return pfad
