"""Smoke-Prüfung: Syntax, Dialog und Exporte mit isolierten Testdaten."""
import ast
import os
from pathlib import Path
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ['KIVY_NO_ARGS'] = '1'
os.environ['KIVY_NO_FILELOG'] = '1'

files = list(ROOT.glob('*.py'))
for folder in ('screens', 'widgets', 'models', 'layouts', 'berichte'):
    files.extend((ROOT / folder).rglob('*.py'))
for path in files:
    ast.parse(path.read_text(encoding='utf-8-sig'), filename=str(path))
print('SYNTAX_OK', len(files))

from pruefe_manuelle_buchungen import ManualTests
case = ManualTests()
case.setUp()
db = case.db
db.save_manual_entry('2026-09-20', 1, 'INCOME', '100', 'Sponsor', 'S-1')
db.save_manual_entry('2026-09-20', 1, 'EXPENSE', '25', 'Raummiete', 'R-1')
db.get_statistic_sale_items = lambda *args: []
db.get_article_sales = lambda *args: []
db.get_category_revenues = lambda *args: []
db.get_entwertet = lambda *args: {'kig_karte': 0, 'gutschein': 0}
db.get_events = lambda: [dict(id=1, name='Fest', entry_type='EVENT', start_date='2026-09-20')]

from berichte import statistik_bericht
from openpyxl import load_workbook
data = statistik_bericht.sammeln(db, (None, None, None, None), 'Integrationstest')
with tempfile.TemporaryDirectory(prefix='kig-statistik-test-') as tmp:
    excel = Path(tmp) / 'test.xlsx'
    pdf = Path(tmp) / 'test.pdf'
    statistik_bericht.excel(data, excel)
    statistik_bericht.pdf(data, pdf)
    book = load_workbook(excel)
    assert 'Manuelle Buchungen' in book.sheetnames
    assert any(cell.value == 'Sponsor' for row in book['Manuelle Buchungen'] for cell in row)
    book.close()
    assert pdf.stat().st_size > 1000
print('EXPORT_OK')

from kivy.config import Config
Config.set('graphics', 'fullscreen', '0')
Config.set('graphics', 'width', '900')
Config.set('graphics', 'height', '700')
from widgets.statistics.manual_entries import ManualEntriesPopup
from kivy.clock import Clock
popup = ManualEntriesPopup(db, (None, None, None, None), lambda: None)
assert len(popup.rows.children) == 2
popup.edit()
Clock.tick()
print('DIALOG_OK')
case.tearDown()
